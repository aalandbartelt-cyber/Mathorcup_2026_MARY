#!/usr/bin/env python3
"""生成 Q1 交付给论文手的模型假设与问题分析独立 DOCX。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[3]
OUT_DIR = REPO / "05_论文" / "Q1_感官评分"


def set_font(style, *, east_asia: str, latin: str, size: float, bold: bool = False):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    color = r_pr.find(qn("w:color"))
    if color is not None:
        color.set(qn("w:val"), "000000")
        color.attrib.pop(qn("w:themeColor"), None)
        color.attrib.pop(qn("w:themeTint"), None)
        color.attrib.pop(qn("w:themeShade"), None)


def make_base() -> Document:
    doc = Document()
    # python-docx 默认写入 <w:zoom w:val="bestFit"/>，严格 OOXML 校验要求
    # zoom 节点包含 w:percent。与已验收的 Q1 正文统一为 100%。
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.attrib.pop(qn("w:val"), None)
        zoom.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.80)
    sec.right_margin = Cm(2.80)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    set_font(normal, east_asia="宋体", latin="Times New Roman", size=12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    body = doc.styles["Body Text"]
    set_font(body, east_asia="宋体", latin="Times New Roman", size=12)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.first_line_indent = Cm(0.74)

    h1 = doc.styles["Heading 1"]
    set_font(h1, east_asia="黑体", latin="Times New Roman", size=14, bold=True)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_font(h2, east_asia="黑体", latin="Times New Roman", size=12, bold=True)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(0)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.keywords = ""
    cp.category = ""
    return doc


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph(style="Body Text")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def build_assumptions() -> Path:
    doc = make_base()
    doc.core_properties.title = "问题一模型假设"
    doc.add_paragraph("模型假设", style="Heading 1")
    add_body(doc, "为简化问题，本文做出以下假设：")

    items = [
        (
            "假设1：",
            "同一化学字段在各啤酒样本中具有可比的单位和测量口径。数据依据是附件1按统一字段给出全部样本，且附件8可逐值核验的70个字段中，37个满足log10变换关系、33个保持原尺度。若发现跨批次量纲漂移，则先利用质控样本完成批次校准，再重新训练和验证模型。",
        ),
        (
            "假设2：",
            "在现有信息下，将每款啤酒视为相互独立的观测。数据依据是beer字段在附件1和附件4中均完整且唯一，而附件未提供酒厂、生产批次或重复配方标识。若后续获得同厂或同批信息，则改用分组交叉验证或层次模型，重新估计泛化误差。",
        ),
        (
            "假设3：",
            "少量缺失不会使整款啤酒失去建模价值，且缺失位置可能携带测量流程信息。数据依据是231个输入字段中仅9列出现缺失，共73个缺失单元、涉及63行，其余化学特征仍可观测。因此在训练折内采用中位数填补并保留缺失指示；若缺失由系统性测量故障主导，则补做多重或迭代填补敏感性分析，并降低相关预测的可信度。",
        ),
        (
            "假设4：",
            "附件4的标准化连续感官值保留各感官维度的相对强弱顺序，可作为多输出回归的监督信号。数据依据是附件3给出0—6等级定义，附件4给出250款啤酒的50个同名标准化感官属性，但未提供完整的评委均值和标准差。若该顺序关系不成立或获得原始评分，则改在原始0—6尺度上建模并废止经验分布映射。",
        ),
        (
            "假设5：",
            "化学组成与感官强度之间的映射在当前商业啤酒样本总体内近似稳定。数据依据是250款样本覆盖22种细分风格，可用于已见风格内的重复分层验证，但没有独立产区或测量平台的外部样本。若应用到完全未见风格、新产区或新平台，则先进行外部验证和校准；若误差明显增大，应重新训练模型。",
        ),
    ]
    for lead, rest in items:
        add_body(doc, lead + rest, bold_lead=lead)

    path = OUT_DIR / "模型假设_Q1.docx"
    doc.save(path)
    return path


def build_analysis() -> Path:
    doc = make_base()
    doc.core_properties.title = "2.1 问题一分析"
    doc.add_paragraph("2.1 问题一分析", style="Heading 2")
    add_body(
        doc,
        "问题一的本质是利用250款啤酒的231项化学与理化特征，建立面向50个Aroma与Flavour感官属性的多输出回归模型，并将标准化连续预测转换为0—6相对强度等级。建模难点主要有三点：一是样本量与输入维数接近，化合物间又存在较强共线性，模型容易过拟合；二是输入包含少量缺失和混合量纲，若在全数据上填补或变换会造成验证信息泄漏；三是附件4给出的是标准化连续值，无法唯一还原原始评委平均分，同时22种细分风格的样本量不均，宏平均指标可能掩盖弱目标和小样本风格的预测失败。",
    )
    add_body(
        doc,
        "针对上述困难，本文先按beer字段对齐样本，在每个训练折内完成中位数填补、缺失指示和必要的特征变换；随后以全局均值、风格收缩均值、PLS2和Ridge构成递进基线，并使用ExtraTrees刻画非线性与特征交互。模型选择采用内层交叉验证，性能评价采用10×3重复风格分层外层交叉验证产生的OOF预测，同时报告逐目标、逐风格误差及配对bootstrap区间。最终保留标准化感官预测，并仅利用训练折经验分布映射到0—6相对强度等级，以保证量表转换不使用验证信息。",
    )

    path = OUT_DIR / "问题分析_Q1.docx"
    doc.save(path)
    return path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for path in (build_assumptions(), build_analysis()):
        print(path)


if __name__ == "__main__":
    main()
