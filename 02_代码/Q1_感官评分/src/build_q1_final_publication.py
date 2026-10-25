"""Build the final Q1 manuscript from the specified Word template and audited tables."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import copy
import json
import re
import shutil
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from build_q1_docx import (
    add_body, add_heading, add_three_line_table, add_equation, add_figure,
    clear_document_body, normalize_ooxml_order, formula_model_relation,
    formula_grade_mapping, omml_sub as sub, omml_sup as sup, omml_bar as bar,
    omml_fraction as frac, omml_nary as sigma, omml_delimiter as delim, omml_accent as accent,
)

ROOT=Path(__file__).resolve().parents[3]
DATA=ROOT/'03_结果/Q1_感官评分/publication_final'
OUT=ROOT/'05_论文/Q1_感官评分'
FIG=ROOT/'04_图表/Q1_感官评分/final_v4'
TEMPLATE=ROOT/'docs/reference/agent初稿/分小问独立文档/Q1.docx'
GROUP_ZH={'Aroma':'香气','Flavour-core':'风味','Taste-other':'其他味觉','Palate':'口感','Overall':'总体喜好'}

ANALYSIS=[
    '问题一要求由化学组成预测多维感官强度。本题可用样本为250款啤酒，输入包含231项化学与理化特征，监督信号为50个标准化连续感官属性。其困难在于输入维数接近样本量、化合物之间相关性较强，各感官家族的可预测性又不相同；附件中的标准化值还不能唯一还原为评委原始0—6分。风格分布不均，单一平均误差也可能掩盖缺陷味及小样本风格的失败。',
    '本文按beer字段对齐附件，在训练折内完成填补与特征变换；以ExtraTrees为非线性基线，将HGB、秩高斯Ridge和PLS2的互补预测纳入非负最小二乘堆叠，并通过感官组向共享权重收缩控制方差。采用10次重复、每次3折的外层验证，基学习器选择与堆叠权重估计均限制在外层训练数据内。评价同时覆盖解释度、绝对误差、排序一致性及弱目标、风格误差，0—6仅按训练折经验分布给出相对强度。',
]
ASSUMPTIONS=[
    '假设1：同一字段在样本之间具有可比的测量含义。附件1采用统一字段，附件8支持识别37个log10字段与33个原尺度字段，但字段尺度可识别不代表每个单元格一致。本文保留附件1给定值；若发现批次或单元格冲突，则单列核验，涉及绝对浓度的解释仅使用逐条确认的观测。',
    '假设2：在没有酒厂、批次及重复配方标识时，将每款啤酒作为一个观测单位。beer字段完整且唯一，支持样本对齐；若后续获得同厂或同批信息，则改用分组交叉验证，并重新估计泛化误差。',
    '假设3：少量缺失不使整款样本失去预测价值，且缺失位置可能包含测量流程信息。输入共73个缺失单元，涉及63行9列，因此采用折内中位数填补并保留指示。若缺失机制明显依赖未观测浓度或系统故障，则补做删失模型或多重插补分析，降低有关预测的可信度。',
    '假设4：附件4中的连续值保留各目标内部的相对强弱顺序，可用于监督学习。附件3给出0—6评分定义，但没有完整的评委中心化、缩放参数，故不假设原始均分可逆恢复；若获得原始评分，改在原始量表建模并重新评价映射。',
    '假设5：当前样本覆盖范围内，不同感官家族可共享部分预测信息。五个家族和重复折内预测支持估计共享及组内权重，固定0.5收缩用于约束小组方差；若外部样本或未见风格出现结构漂移，则重新校准、分组验证或拒绝外推。',
]

def base_document(title):
    # Start from a physical copy, then replace content blocks inside that copy.
    working=ROOT/'outputs/qa/q1_final/template_working_copy.docx'
    working.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(TEMPLATE,working)
    doc=Document(working)
    proto=copy.deepcopy(doc.tables[0]._tbl)
    if proto.find(qn('w:tblGrid')) is None:
        grid=OxmlElement('w:tblGrid')
        for width in [Cm(1.02),Cm(13.20),Cm(1.02)]:
            node=OxmlElement('w:gridCol');node.set(qn('w:w'),str(width.twips));grid.append(node)
        proto.insert(1,grid)
    clear_document_body(doc)
    cp=doc.core_properties
    cp.title=title;cp.subject='啤酒化学组成与多维感官预测'
    for k in ['author','last_modified_by','comments','keywords','category','identifier']:
        setattr(cp,k,'')
    # Preserve the template's section settings; final-paper assembly controls physical page geometry.
    sec=doc.sections[0]
    return doc,proto

def clean_unused_parts(path):
    """Remove unused sample images, notes and personal properties inherited from the template."""
    with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    with ZipFile(TEMPLATE) as z:
        # Keep the original style definitions byte-for-byte, including linked styles.
        parts['word/styles.xml']=z.read('word/styles.xml')
    ns={'r':'http://schemas.openxmlformats.org/package/2006/relationships'}
    root=etree.fromstring(parts['word/document.xml'])
    used=set(root.xpath('//@r:embed | //@r:link | //@r:id',namespaces={'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}))
    rel=etree.fromstring(parts['word/_rels/document.xml.rels'])
    keep_media=set()
    for node in list(rel):
        target=node.get('Target','');typ=node.get('Type','')
        if typ.endswith('/image'):
            if node.get('Id') in used:keep_media.add('word/'+target)
            else:rel.remove(node)
        elif any(typ.endswith('/'+s) for s in ['comments','footnotes','endnotes','customXml']):
            rel.remove(node)
    parts['word/_rels/document.xml.rels']=etree.tostring(rel,xml_declaration=True,encoding='UTF-8',standalone=True)
    removed=[]
    for name in list(parts):
        if (name.startswith('word/media/') and name not in keep_media) or name.startswith('customXml/') or name in ['word/footnotes.xml','word/endnotes.xml','word/comments.xml','docProps/custom.xml']:
            removed.append(name);del parts[name]
    # Remove dangling part declarations and package-level custom-property references.
    for name in ['[Content_Types].xml','_rels/.rels']:
        tree=etree.fromstring(parts[name])
        for node in list(tree):
            target=node.get('PartName','').lstrip('/') or node.get('Target','')
            if target in removed:tree.remove(node)
        parts[name]=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)

def save(doc,path):
    # Style names alone do not capture the template's direct paragraph/run formatting.
    template_doc=Document(TEMPLATE)
    prototypes={name:next(p for p in template_doc.paragraphs if p.style.name==name and p.text)
                for name in ['Heading 1','Heading 2','Heading 3','Body Text']}
    for p in doc.paragraphs:
        if not p.text or p.style.name not in prototypes:continue
        if any(run.font.name=='Consolas' for run in p.runs):continue
        if re.match(r'^(表\d|图\d|附图S)',p.text) and p.alignment==1:continue
        proto=prototypes[p.style.name]
        old=p._p.find(qn('w:pPr'))
        if old is not None:p._p.remove(old)
        p._p.insert(0,copy.deepcopy(proto._p.pPr))
        original=next(r for r in proto.runs if r.text.strip())._r.rPr
        for run in p.runs:
            old=run._r.find(qn('w:rPr'))
            if old is not None:run._r.remove(old)
            run._r.insert(0,copy.deepcopy(original))
    normalize_ooxml_order(doc)
    sequence=['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize',
              'tblW','jc','tblCellSpacing','tblInd','tblBorders','shd','tblLayout','tblCellMar',
              'tblLook','tblCaption','tblDescription','tblPrChange']
    order={qn('w:'+name):i for i,name in enumerate(sequence)}
    for table in doc.tables:
        pr=table._tbl.tblPr
        children=list(pr)
        for node in children:pr.remove(node)
        for node in sorted(children,key=lambda n:order.get(n.tag,99)):pr.append(node)
        if table.rows:
            for grid,cell in zip(table._tbl.tblGrid,table.rows[0].cells):
                if cell.width is not None:grid.set(qn('w:w'),str(cell.width.twips))
    doc.save(path);clean_unused_parts(path)
    return path

def formula_objective():
    inner=[sub(['y'],['ij']),' − ',sub([bar(['y'])],['j']),' − ',
           sigma([sub(['w'],['m']),delim([sub(['z'],['ijm']),' − ',sub([bar(['z'])],['jm'])])],['m=1'],['4'])]
    loss=sup([delim([frac(inner,[sub(['s'],['j'])])])],['2'])
    return [sup([sub(['w'],['G'])],['*']),' = ',sub(['arg min'],['w≥0']),
            sigma([sigma([loss],['i∈T'])],['j∈G'])]

def formula_shrink():
    return [sub([accent(['w'],char='̃')],['g']),' = (1−λ)',sup([sub(['w'],['g'])],['*']),
            ' + λ',sup([sub(['w'],['全部'])],['*']),',    λ = 0.5']

def formula_intercept():
    return [sub(['b'],['j']),' = ',sub([bar(['y'])],['j']),' − ',
            sigma([sub([accent(['w'],char='̃')],['g(j),m']),sub([bar(['z'])],['jm'])],['m=1'],['4'])]

def formula_prediction():
    return [sub([accent(['y'])],['ij']),' = ',sub(['b'],['j']),' + ',
            sigma([sub([accent(['w'],char='̃')],['g(j),m']),sup([sub(['f'],['m'])],['(T)']),delim([sub(['x'],['i'])])],['m=1'],['4'])]

def formula_macro_r2():
    numerator=sigma([sup([delim([sub(['y'],['ij']),'−',sub([accent(['y'])],['ij'])])],['2'])],['i=1'],['n'])
    denominator=sigma([sup([delim([sub(['y'],['ij']),'−',sub([bar(['y'])],['j'])])],['2'])],['i=1'],['n'])
    return [sub([sup(['R'],['2'])],['宏']),' = ',frac(['1'],['J']),sigma([delim(['1−',frac([numerator],[denominator])],begin='[',end=']')],['j=1'],['J'])]

def formula_macro_mae():
    return [sub(['MAE'],['宏']),' = ',frac(['1'],['nJ']),sigma([sigma([delim([sub(['y'],['ij']),'−',sub([accent(['y'])],['ij'])],begin='|',end='|')],['i=1'],['n'])],['j=1'],['J'])]

def build():
    OUT.mkdir(parents=True,exist_ok=True)
    metrics=pd.read_csv(DATA/'model_summary.csv');summary=json.loads((DATA/'final_model_summary.json').read_text())
    target=pd.read_csv(DATA/'target_metrics_final.csv')
    style=pd.read_csv(DATA/'style_metrics_final.csv')
    ablation=pd.read_csv(DATA/'weight_structure_ablation.csv')
    weights=pd.read_csv(DATA/'stacking_weights_summary.csv')
    bootstrap=pd.read_csv(DATA/'paired_beer_bootstrap.csv')
    doc,proto=base_document('问题一模型建立与求解')
    add_heading(doc,'问题一模型建立与求解',1)
    add_heading(doc,'问题一分析',2)
    for text in ANALYSIS:add_body(doc,text)
    add_heading(doc,'问题一模型',2)
    add_heading(doc,'数据口径与缺失值插补',3)
    add_body(doc,'主键使用beer，beer_id存在精度问题，不参与连接或建模；风格字段仅用于分层验证及分组误差报告，不作为最终模型输入。表1给出预处理步骤与维度变化。所有需估计的填补、变换及映射参数只由相应训练折计算。附件1的给定尺度保持不变，已确认的log10字段不重复取对数；其中负值表示原浓度小于1。')
    add_three_line_table(doc,'表1 数据预处理步骤与维度变化',['步骤','处理与范围','维度变化','理由'],[
        ('读取核对','附件1与附件4','250×234；250×53','核对主键、目标与字段类型'),
        ('样本对齐','按beer一一连接','250+250→250行','避免标识精度导致错配'),
        ('字段拆分','剔除三列元数据','X：250×231；Y：250×50','排除身份与风格标签入模'),
        ('尺度保留','不重复取log10、不全局截尾','250行、231输入不变','保留实际差异与极端观测'),
        ('折内插补','训练中位数与缺失指示','231→231+k，0≤k≤9','73单元、63行、9列缺失'),
        ('分支变换','线性分支折内秩高斯变换','样本及目标数不变','减轻偏态、共线与量纲影响'),
        ('相对映射','仅用训练折经验分布','50维连续值＋50维相对等级','不声称恢复原始评委均分'),
    ],widths=[Cm(2),Cm(4.0),Cm(4),Cm(5.6)],font_size=8.3)
    add_body(doc,'全部250行均保留，不按全数据分位数删除异常样本。37个log10字段及33个原尺度字段可由附件8识别，但这不意味着两附件所有对应单元格一致；发现冲突时保留输入原值并单列核验，不静默替换。缺失机制及跨批次可比性属于模型假设，需由后续外部样本检验。')
    add_heading(doc,'连续监督信号与相对标签',3)
    add_body(doc,'以标准化感官向量作为监督信号，多输出关系定义为式(1)。每款啤酒的化学特征向量记为x，感官向量记为y；误差项包含未测化合物、评委差异及随机扰动。')
    add_equation(doc,formula_model_relation(),1,proto)
    add_body(doc,'附件4包含标准化后汇总值，缺少完整的评委中心与尺度参数，因此从它到原始0—6均分并非唯一。连续尺度承担模型训练与误差评价，相对等级仅承担展示；两者分开报告。')
    add_heading(doc,'基学习器与特征构造',3)
    add_body(doc,'基学习器由ExtraTrees、直方图梯度提升HGB、秩高斯Ridge和秩高斯PLS2组成。ExtraTrees刻画随机化非线性与交互[2]，HGB补充局部阈值响应，Ridge缓和共线，PLS2提供多输出共享的低秩表示。树分支保留附件尺度；线性分支在各训练折中填补后进行秩高斯变换，并将缺失指示单独并入。表2列出完整候选范围。')
    add_three_line_table(doc,'表2 基学习器与折内候选参数',['学习器','候选范围','数据处理'],[
        ('ExtraTrees','300棵；特征比例0.33/0.60/1.00；叶节点1/2/4/8；不限深度','中位数填补＋缺失指示'),
        ('HGB','叶节点上限7/15；最小叶样本20；学习率0.05；100轮；L2=1','中位数填补＋缺失指示'),
        ('Ridge','正则强度0.1/1/10/100/1000','折内秩高斯变换与缩放'),
        ('PLS2','潜变量数2/5/10/15/20','折内秩高斯变换与多输出拟合'),
    ],widths=[Cm(2.3),Cm(9.2),Cm(4.1)],font_size=8.3)
    add_body(doc,'内层以宏平均MAE最小选择每个学习器的参数；相同时依次比较R²、Spearman和两项七档代理指标。候选范围固定，外层验证数据不参与参数、权重或截距估计。HGB候选也可能得到相同预测，本文保留固定候选范围，不据此扩大搜索。')
    add_heading(doc,'感官组收缩非负堆叠模型',3)
    add_body(doc,'堆叠利用训练数据内部的交叉验证预测估计组合系数，其非负约束可减少相关预测器之间的相互抵消[1]。设T为当前外层训练样本，z为各基学习器在T上经内层3折得到的折外预测，m为四个基学习器的索引。对目标集合G，分别求解式(2)；G可取全部目标或某一感官家族。')
    add_equation(doc,formula_objective(),2,proto)
    add_body(doc,'式中，y的横线为训练目标均值，z的横线为对应内层折外预测均值，s为训练目标的标准差；零方差时取1。标准差归一化使高波动目标不独占权重拟合。按香气15维、风味15维、其他味觉15维、口感4维、总体喜好1维分组；组别由字段定义预先确定，与啤酒风格不同。')
    add_body(doc,'为降低小感官组的权重方差，将组内解向全目标共享解收缩，定义式(3)。固定λ=0.5，不使用外层指标选择收缩强度。')
    add_equation(doc,formula_shrink(),3,proto)
    add_body(doc,'非负约束只要求各权重不小于0，不要求权重之和为1。每个目标保留自由截距，按训练均值恢复为式(4)，因此截距可为负值，不能与权重一同约束为非负。')
    add_equation(doc,formula_intercept(),4,proto)
    add_body(doc,'式中g(j)为目标j所属感官家族。各基学习器随后在完整外层训练集T上重拟合，对验证样本的最终连续预测由式(5)给出。')
    add_equation(doc,formula_prediction(),5,proto)
    add_body(doc,'表3汇总30个外层训练折的权重，表示不同预测器在本数据中的组合方式，不作因果贡献比例解释。权重总和允许超过1，自由截距同时校正中心偏移。')
    wt=weights.pivot(index='target_group',columns='base_model',values='weight_mean')
    add_three_line_table(doc,'表3 感官组收缩堆叠的平均非负权重',['感官家族','ExtraTrees','HGB','Ridge','PLS2'],[
        (GROUP_ZH[g],*[f'{wt.loc[g,b]:.3f}' for b in ['ExtraTrees','HGB','Ridge（秩高斯）','PLS2（秩高斯）']])
        for g in GROUP_ZH],widths=[Cm(3.6),Cm(3),Cm(3),Cm(3),Cm(3)],font_size=9)
    add_heading(doc,'0—6相对强度映射',3)
    add_body(doc,'仅利用当前外层训练集中目标j的经验分布，将验证预测单调映射至0—6区间，定义为式(6)。其中F为训练目标的经验分布函数，指示函数统计不大于给定预测值的训练观测比例。')
    add_equation(doc,formula_grade_mapping(),6,proto)
    add_body(doc,'先逐折映射，再对10次重复的等级取平均，不能先汇总预测再用全体目标拟合映射。该等级保持维度内部的相对强弱次序，不能等同于原始评委平均分。附图S1(b)展示固定训练折的映射曲线。')
    add_heading(doc,'问题一求解',2)
    add_heading(doc,'求解步骤与交叉验证',3)
    add_body(doc,'使用固定的10次重复、每次3折外层划分，尽量平衡细分风格；极小风格不强求每折均出现。每次重复中，每款啤酒恰好进入一个验证折。对每个外层训练集，依次进行内层3折参数选择、内层折外预测构造、共享及组内NNLS拟合、0.5权重收缩与自由截距恢复，最后重拟合基学习器并预测外层验证集。随机种子以20261023为起点，并由重复、折及学习器索引确定。')
    add_body(doc,'对每个感官目标先单独计算指标，再对50个目标等权平均。式(7)、式(8)给出宏平均R²与MAE；n为本次评价样本数，J=50，评价均值仅用于指标定义，不用于训练。RMSE同样先按目标计算再平均。')
    add_equation(doc,formula_macro_r2(),7,proto);add_equation(doc,formula_macro_mae(),8,proto)
    add_body(doc,'Spearman衡量连续值排序关系。七档分位代理QWK与相邻等级一致率分别将观测、预测按各自样本分位数分为七档，再计算一致性；它们是固定的事后排序诊断，不进入相对等级推断，也不是原始0—6评分准确率。以下主比较使用10个完整重复指标的均值与标准差；平均折外预测则先对每款啤酒的10次预测取平均，二者单列。')
    add_heading(doc,'结果分析与模型对比',3)
    add_body(doc,'表4比较非线性基线及三种融合方法与最终模型，各模型使用相同外层划分，图1进一步展示重复间波动。最终模型提高解释度与连续预测精度，但在部分排序指标上没有取得最高值。')
    rows=[]
    for _,r in metrics.iterrows():
        rows.append((r.display_name,f'{r.r2_mean:.4f}±{r.r2_sd:.4f}',f'{r.mae_mean:.4f}±{r.mae_sd:.4f}',f'{r.spearman_mean:.4f}',f'{r.qwk_proxy_mean:.4f}',f'{100*r.acc_pm1_proxy_mean:.2f}%'))
    add_three_line_table(doc,'表4 相同外层划分下的五项预测指标',['方法','R²（均值±SD）','MAE（均值±SD）','Spearman','代理QWK','相邻等级一致率'],rows,widths=[Cm(3.1),Cm(3.3),Cm(3.3),Cm(2),Cm(1.8),Cm(2.1)],font_size=7.8)
    add_body(doc,'感官组收缩NNLS的重复均值R²为0.2564，MAE为0.3361，Spearman为0.4302，代理QWK为0.4183，相邻等级一致率为55.42%；五项均在10次重复中优于ExtraTrees基线。重复均值RMSE为0.4714。连续—有序融合的MAE为0.3356、Spearman为0.4314、代理QWK为0.4194、相邻等级一致率为55.52%，均略好于最终模型。选择最终模型依据连续解释度、平均折外精度及稳健性综合权衡，不将其表述为所有指标同时最优。')
    add_figure(doc,FIG/'图1_模型对比.png','图1 基线、融合与最终模型的重复折外性能比较（纵轴为局部范围）')
    mo=summary['mean_oof']
    add_body(doc,f"平均折外预测的宏平均R²为{mo['r2']:.4f}、MAE为{mo['mae']:.4f}、RMSE为{mo['rmse']:.4f}、Spearman为{mo['spearman']:.4f}，代理QWK为{mo['qwk_proxy']:.4f}、相邻等级一致率为{100*mo['acc_pm1_proxy']:.2f}%。ExtraTrees基线对应R²为0.2502、MAE为0.3389。平均预测降低了随机划分造成的波动，但这些指标不替代上文10次重复的均值；逐目标收缩融合的平均折外Spearman为0.4511，仍高于最终模型。")
    add_body(doc,f"图2按感官家族完整展示50个目标，避免只挑高分目标。目前{int((target.r2>0).sum())}个目标的平均折外R²为正。酸度、酒精感等信号较强，花香酯、陈旧酒花味及部分总体感受信号较弱。每个目标的解释度与秩相关不能互相替代，尤其在预测幅度收缩时，排序尚可仍可能伴随较低R²。")
    add_figure(doc,FIG/'图2_逐目标预测性能.png','图2 最终模型50个感官目标的平均折外R²与Spearman')
    add_body(doc,'图3沿用事先固定的酸度、4VG和陈旧酒花味三个代表目标，分别检视较强、中等及较弱信号。虚线为理想预测，实线与阴影为观测—预测OLS拟合及均值置信带，不是单个新样本的预测区间。弱目标的预测仍向均值收缩，说明共享建模不能补足未测化学信息。')
    add_figure(doc,FIG/'图3_OOF观测预测.png','图3 固定代表感官目标的平均折外观测—预测关系')
    add_heading(doc,'模型检验与稳健性',3)
    add_body(doc,'输入中排除标识与风格标签，填补、秩变换、基学习器选择、堆叠与映射均在相应训练范围内进行。附件2、5、6的全量相关矩阵不用于本问筛选或调参。数据附带的标准化监督信号沿用官方定义；没有原始评委记录时，不能审计或消除其上游标准化过程可能带来的依赖。')
    add_body(doc,'为检验弱目标，以基线重复平均R²预先固定最低的13个目标，另固定10个难目标，不能按最终模型重新挑选。尾部13目标的平均R²从0.0411增至0.0639，MAE从0.3277降至0.3227；10个难目标平均ΔR²为+0.0081、ΔMAE为−0.0032。相较逐目标收缩融合，尾部MAE仍略高，因此不声称弱目标普遍最优。图4同时展示风格误差和难目标的正负变化。')
    add_figure(doc,FIG/'图4_风格与难目标稳健性.png','图4 平均折外风格误差与预登记难目标的重复平均R²变化')
    add_body(doc,'22个风格中，20个的重复平均MAE低于基线；最差风格Brut由0.5512降至0.5348。佛兰德老棕艾尔与Tripel分别增加0.0019和0.0011。总体喜好和反式-2-壬烯醛的R²分别下降0.0144和0.0101，MAE分别增加0.0042和0.0010；香气、风味的热带酒花维度也有小幅下降。风格内样本量较小，这些描述不构成未见风格的外推保证。')
    add_body(doc,'进一步在固定平均折外预测上，以啤酒为配对单位进行20,000次自助重采样。每款啤酒先对50个目标取绝对误差平均；秩误差则对各目标的观测和预测分别排名，以样本数减1归一化，再求绝对差。表5中的差值均为最终模型减基线，负值表示误差较低。')
    add_three_line_table(doc,'表5 最终模型相对基线的配对自助重采样结果',['指标','基线','最终模型','差值','95%百分位区间'],[
        ('MAE',f"{bootstrap.iloc[0]['baseline']:.6f}",f"{bootstrap.iloc[0]['candidate']:.6f}",f"{bootstrap.iloc[0]['delta_candidate_minus_baseline']:.6f}",f"[{bootstrap.iloc[0]['ci_low']:.6f}, {bootstrap.iloc[0]['ci_high']:.6f}]"),
        ('归一化绝对秩误差',f"{bootstrap.iloc[1]['baseline']:.6f}",f"{bootstrap.iloc[1]['candidate']:.6f}",f"{bootstrap.iloc[1]['delta_candidate_minus_baseline']:.6f}",f"[{bootstrap.iloc[1]['ci_low']:.6f}, {bootstrap.iloc[1]['ci_high']:.6f}]"),
    ],widths=[Cm(3.2),Cm(2.5),Cm(2.5),Cm(2.5),Cm(4.9)],font_size=8.3)
    add_body(doc,'两个区间均位于0以下，支持当前固定预测下的平均改善。该自助区间不包含重训或模型选择的不确定性，10次划分也共享样本，不作为10个独立研究处理。附图S1(a)继续检查残差的极端区间偏差与方差变化。')
    add_body(doc,'表6比较预先设置的三种权重共享结构；共享权重反映全局共性，逐目标权重提供更细自由度，感官组结构处于二者之间。三种模型均固定0.5收缩，表内比较仅解释权重结构差异，不是对收缩系数的敏感性曲线。')
    add_three_line_table(doc,'表6 权重共享结构消融',['结构','R²','MAE','Spearman','代理QWK','相邻等级一致率'],[
        (r['model'],f"{r['r2']:.4f}",f"{r['mae']:.4f}",f"{r['spearman']:.4f}",f"{r['qwk_proxy']:.4f}",f"{100*r['acc_pm1_proxy']:.2f}%") for _,r in ablation.iterrows()
    ],widths=[Cm(4.1),Cm(2.2),Cm(2.2),Cm(2.4),Cm(2.2),Cm(2.5)],font_size=8.3)
    add_body(doc,'感官组收缩在此比较中取得较高解释度与较低MAE；逐目标收缩的相邻等级一致率为55.432%，略高于感官组的55.423%，说明更细的权重并未带来各项指标一致改善。现有结果只支持固定结构下的比较；其他收缩强度、跨批次数据和原始0—6标签的检验仍待补充。')
    add_heading(doc,'结果解释与模型评价',3)
    add_body(doc,'最终模型利用树模型、线性收缩与低秩多输出的误差互补，在严格外层验证中提高整体连续解释度，固定尾部和最差风格也优于基线。NNLS权重保留可解释的非负组合，自由截距与组间收缩控制幅度和小组方差，适合作为后续感官画像的预测接口。')
    add_body(doc,'其限制有三点：少数难目标的化学信号仍不足，整体改善不能替代逐维可靠性；七档排序代理及相对强度映射均不能当作原始评委评分；重复验证、自助区间和标准化数据只支持当前商业啤酒样本范围，不能证明未见风格或合成化学谱上的真实感官安全。问题二的关联证据与问题三的候选设计仍须分别验证。')
    add_heading(doc,'参考文献',2)
    add_body(doc,'[1] Breiman L. Stacked regressions. Machine Learning, 1996, 24: 49–64. DOI: 10.1023/A:1018046112532.')
    add_body(doc,'[2] Geurts P, Ernst D, Wehenkel L. Extremely randomized trees. Machine Learning, 2006, 63: 3–42. DOI: 10.1007/s10994-006-6226-1.')
    add_heading(doc,'附图',2)
    add_body(doc,'附图S1合并残差诊断与训练折经验分布映射。残差图使用与图3相同的固定目标；映射示例仅使用首个重复的首个外层训练折，不使用验证观测拟合经验分布。程序及完整逐维表另附。')
    add_figure(doc,FIG/'附图S1_残差诊断与量表映射.png','附图S1 固定目标残差与训练折经验分布的相对量表映射')
    path=save(doc,OUT/'第一问_论文正文_v5.docx')
    # Separate reusable sections share the same template styles.
    for title,texts,filename in [('模型假设',['为简化问题，本文做出以下假设：',*ASSUMPTIONS],'模型假设_Q1_v3.docx'),('问题一分析',ANALYSIS,'问题分析_Q1_v3.docx')]:
        d,_=base_document(title);add_heading(d,title,1)
        for text in texts:add_body(d,text)
        save(d,OUT/filename)
    abstracts=[f"针对化学特征到多维感官评分的预测，建立感官组收缩非负堆叠模型，以ExtraTrees、HGB、秩高斯Ridge及PLS2为基学习器，在10×3外层验证中以折内预测估计非负权重和自由截距。重复均值R²为0.2564、MAE为0.3361；平均折外预测R²为{mo['r2']:.4f}、MAE为{mo['mae']:.4f}。基线固定尾部13目标的平均R²由0.0411增至0.0639，配对自助重采样的MAE差95%区间为[−0.0095,−0.0053]。部分排序指标和难目标仍逊于其他方法，0—6输出仅作相对强度。"]
    (OUT/'摘要结论与模型评价.md').write_text('# 问题一摘要用语\n\n'+abstracts[0]+'\n\n# 结论与评价\n\n'+ '\n\n'.join(p.text for p in doc.paragraphs if p.text.startswith(('最终模型利用','其限制有三点','感官组收缩NNLS的重复')))+'\n')
    (OUT/'01_问题分析.md').write_text('# 问题一分析\n\n'+'\n\n'.join(ANALYSIS)+'\n')
    (OUT/'02_模型假设与口径.md').write_text('# 模型假设\n\n为简化问题，本文做出以下假设：\n\n'+'\n\n'.join(ASSUMPTIONS)+'\n')
    return path

if __name__=='__main__':
    print(build())
    import subprocess,sys
    subprocess.run([sys.executable,str(ROOT/'scripts/export_paper_review.py'),'--question','Q1_感官评分'],check=True)
