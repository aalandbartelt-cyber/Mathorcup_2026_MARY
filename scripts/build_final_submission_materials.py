#!/usr/bin/env python3
"""Build the single result workbook and the structured supporting-material tree.

This script intentionally excludes the official raw attachments, model caches,
checkpoints, bytecode, weights, temporary logs, and version-control metadata.
Existing delivery directories are moved to a timestamped backup before the new
validated directories are installed.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


ROOT = Path(__file__).resolve().parents[1]
DELIVERY_ROOT = ROOT / "06_交付"
RESULT_TARGET = DELIVERY_ROOT / "02_结果文件"
SUPPORT_TARGET = DELIVERY_ROOT / "04_支撑材料"
BACKUP_ROOT = ROOT / "_archive/20261028_最终交付整理/目录备份"
TIMESTAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

RESULT_BOOK_NAME = "A2601605_三问结果.xlsx"
TEXT_EXTENSIONS = {
    ".py", ".r", ".R", ".md", ".txt", ".yaml", ".yml", ".json",
    ".csv", ".tsv", ".sh", ".ps1",
}
SKIP_PARTS = {
    "__pycache__", ".pytest_cache", ".mypy_cache", ".ruff_cache", ".git",
    ".idea", ".vscode",
}
SKIP_SUFFIXES = {".pyc", ".pyo", ".joblib", ".pkl", ".pickle"}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def require(path: Path) -> Path:
    if not path.exists():
        raise FileNotFoundError(path)
    return path


def sanitize_text(text: str) -> str:
    """Remove local user paths while keeping repo-root execution semantics."""
    replacements = {
        "/home/mary/miniconda3/envs/mathorcup/bin/python": "python",
        "/home/mary/miniconda3/envs/r/bin/Rscript": "Rscript",
        "/home/mary/miniconda3": "CONDA_ROOT",
        "/home/mary/mathorcup": ".",
        "/mnt/d/上海大学__": "EXTERNAL_REFERENCE_ROOT",
        "/mnt/d/上海大学": "EXTERNAL_REFERENCE_ROOT",
        r"D:\上海大学": "EXTERNAL_REFERENCE_ROOT",
        r"C:\ProgramData": "SYSTEM_PROGRAM_ROOT",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def copy_file(src: Path, dst: Path, *, sanitize: bool = False) -> None:
    require(src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    if sanitize and src.suffix in TEXT_EXTENSIONS:
        try:
            text = src.read_text(encoding="utf-8-sig")
            dst.write_text(sanitize_text(text), encoding="utf-8")
            shutil.copystat(src, dst)
            return
        except UnicodeDecodeError:
            pass
    shutil.copy2(src, dst)


def copy_tree(src: Path, dst: Path, *, sanitize: bool = False) -> list[Path]:
    require(src)
    copied: list[Path] = []
    for path in sorted(src.rglob("*")):
        if not path.is_file():
            continue
        rel = path.relative_to(src)
        if any(part in SKIP_PARTS for part in rel.parts):
            continue
        if path.suffix in SKIP_SUFFIXES or path.name in {".DS_Store", ".gitkeep"}:
            continue
        out = dst / rel
        copy_file(path, out, sanitize=sanitize)
        copied.append(out)
    return copied


def install_directory(staging: Path, target: Path) -> Path | None:
    """Install staging atomically and move any existing target to a backup."""
    backup = None
    if target.exists():
        BACKUP_ROOT.mkdir(parents=True, exist_ok=True)
        backup = BACKUP_ROOT / f"{target.name}_{TIMESTAMP}"
        suffix = 1
        while backup.exists():
            backup = BACKUP_ROOT / f"{target.name}_{TIMESTAMP}_{suffix}"
            suffix += 1
        target.rename(backup)
    staging.rename(target)
    return backup


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out = out.replace([np.inf, -np.inf], np.nan)
    return out.astype(object).where(pd.notna(out), None)


def add_dataframe_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    if len(name) > 31:
        raise ValueError(f"Excel sheet name too long: {name}")
    ws = wb.create_sheet(name)
    df = clean_dataframe(df)
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=9)
    thin = Side(style="thin", color="B7B7B7")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    sample_n = min(ws.max_row, 301)
    for col_idx in range(1, ws.max_column + 1):
        values = [ws.cell(r, col_idx).value for r in range(1, sample_n + 1)]
        width = max((len(str(v)) for v in values if v is not None), default=8) + 2
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width, 10), 34)
    ws.row_dimensions[1].height = 32


def add_key_value_sheet(wb: Workbook, name: str, rows: list[tuple[str, str]]) -> None:
    add_dataframe_sheet(wb, name, pd.DataFrame(rows, columns=["项目", "说明"]))
    ws = wb[name]
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows(min_row=2, max_col=2):
        row[1].alignment = Alignment(vertical="top", wrap_text=True)


def flatten_json(obj: object, prefix: str = "") -> list[tuple[str, object]]:
    rows: list[tuple[str, object]] = []
    if isinstance(obj, dict):
        for key, value in obj.items():
            new_prefix = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(flatten_json(value, new_prefix))
    elif isinstance(obj, list):
        rows.append((prefix, json.dumps(obj, ensure_ascii=False)))
    else:
        rows.append((prefix, obj))
    return rows


def build_result_workbook(out_path: Path) -> dict[str, object]:
    source_root = ROOT / "05_论文/总论文/数据与结果_v2"
    q1_path = require(source_root / "问题一_250啤酒50维相对评分.csv.gz")
    q2_path = require(source_root / "问题二_50维Top3与证据.csv")
    q3_spectra_path = require(source_root / "问题三_900条生成化学谱.csv.gz")
    q3_profile_path = require(source_root / "问题三_完整风格画像.csv")
    q3_migration_path = require(source_root / "迁移状态全集.csv")
    target_map_path = require(source_root / "感官目标中文与模板列对照.csv")
    q1_summary_path = require(ROOT / "03_结果/Q1_感官评分/publication_final/final_model_summary.json")
    q1_target_path = require(ROOT / "03_结果/Q1_感官评分/publication_final/target_metrics_final.csv")

    target_map = pd.read_csv(target_map_path)
    target_lookup = target_map[
        ["registered_target_code", "target_name_zh", "target_name_en", "target_group"]
    ].drop_duplicates("registered_target_code")

    q1 = pd.read_csv(q1_path).merge(
        target_lookup,
        left_on="target_code",
        right_on="registered_target_code",
        how="left",
        validate="many_to_one",
    )
    q1 = q1[
        [
            "beer", "tasting_category_fine", "target_code", "target_name_zh",
            "target_name_en", "target_group", "y_pred_mean", "prediction_sd",
            "grade_0_6", "grade_sd", "relative_grade_integer",
        ]
    ].rename(columns={
        "beer": "啤酒编号",
        "tasting_category_fine": "细分风格",
        "target_code": "感官维度代码",
        "target_name_zh": "感官维度中文名",
        "target_name_en": "感官维度英文名",
        "target_group": "感官家族",
        "y_pred_mean": "标准化感官预测均值",
        "prediction_sd": "重复预测标准差",
        "grade_0_6": "相对强度等级_连续0至6",
        "grade_sd": "相对强度等级标准差",
        "relative_grade_integer": "相对强度等级_整数0至6",
    })
    if q1.shape[0] != 250 * 50 or q1["啤酒编号"].nunique() != 250:
        raise AssertionError(f"Unexpected Q1 shape: {q1.shape}")

    q1_target = pd.read_csv(q1_target_path)
    q1_target = q1_target.merge(
        target_lookup,
        left_on="target_code",
        right_on="registered_target_code",
        how="left",
        validate="one_to_one",
    )
    q1_target = q1_target.drop(columns=["registered_target_code"], errors="ignore")
    q1_summary = json.loads(q1_summary_path.read_text(encoding="utf-8"))
    q1_summary_df = pd.DataFrame(flatten_json(q1_summary), columns=["指标", "数值或说明"])

    q2 = pd.read_csv(q2_path).rename(columns={
        "target_code": "感官维度代码",
        "target_name_en": "感官维度英文名",
        "consensus_rank": "维度内名次",
        "feature_code": "化学字段代码",
        "compound_name": "化合物名称",
        "bootstrap_selection_frequency": "Bootstrap入选频率",
        "leave_style_top3_frequency": "留一风格Top3频率",
        "chemical_evidence_grade": "化学证据等级",
        "oav_median": "气味活性值中位数",
        "claim_scope": "结论边界",
    })
    if q2.shape[0] != 150 or set(q2["维度内名次"].unique()) != {1, 2, 3}:
        raise AssertionError(f"Unexpected Q2 Top-3 table: {q2.shape}")

    q3_profile = pd.read_csv(q3_profile_path).merge(
        target_lookup,
        left_on="target_code",
        right_on="registered_target_code",
        how="left",
        validate="many_to_one",
    )
    q3_profile = q3_profile.drop(columns=["registered_target_code"], errors="ignore")
    q3_migration = pd.read_csv(q3_migration_path)
    q3_spectra = pd.read_csv(q3_spectra_path)
    if q3_profile.shape[0] != 3 * 50 or q3_spectra.shape[0] != 3 * 300:
        raise AssertionError(
            f"Unexpected Q3 output shapes: profile={q3_profile.shape}, spectra={q3_spectra.shape}"
        )
    if q3_migration.shape[0] != 33:
        raise AssertionError(f"Unexpected Q3 migration count: {q3_migration.shape}")

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "参赛队"
    wb.properties.lastModifiedBy = "参赛队"
    wb.properties.title = "MathorCup 2026 A赛道三问结果"
    wb.properties.subject = "感官预测、关键化合物与风格生成迁移结果"
    wb.properties.description = "唯一结果工作簿；各工作表口径见说明与字段说明。"

    instructions = [
        ("文件用途", "本工作簿汇总问题一、问题二和问题三的正式表格结果；结果目录仅保留本文件。"),
        ("问题一口径", "250款啤酒×50个感官维度。标准化预测为10×3外层重复折外预测均值；0-6数值为训练折经验分布上的相对强度映射，不解释为可唯一恢复的原始评委绝对分。"),
        ("问题二口径", "每个感官维度列出3个稳健关键预测关联，共150条；Bootstrap、留一风格和化学证据用于稳定性与合理性核验，不作因果解释。独立模型归因不改变Top-3排序。"),
        ("问题三口径", "给出Tripel、Kriek、Lambic各300条生成化学谱、3×50维感官画像以及33组跨风格迁移审计状态。生成与迁移为统计设计结果，不等同于真实酿造验证或物理安全认证。"),
        ("缺陷约束边界", "DMS、双乙酰、4VG、反-2-壬烯醛等仅在已知且可比的直接浓度或代理口径下审计；未知检出限、源缺失及代理未接入项均不得宣称已验证安全。"),
        ("数据来源", "仅由仓库内已冻结的派生结果生成；本工作簿不包含赛题官方原始附件。"),
    ]
    add_key_value_sheet(wb, "说明", instructions)
    add_dataframe_sheet(wb, "Q1_相对评分", q1)
    add_dataframe_sheet(wb, "Q1_总体指标", q1_summary_df)
    add_dataframe_sheet(wb, "Q1_逐目标指标", q1_target)
    add_dataframe_sheet(wb, "Q2_Top3", q2)
    add_dataframe_sheet(wb, "Q3_生成画像", q3_profile)
    add_dataframe_sheet(wb, "Q3_迁移状态", q3_migration)
    add_dataframe_sheet(wb, "Q3_生成化学谱", q3_spectra)

    field_rows = [
        ("Q1_相对评分", "相对强度等级_连续0至6", "训练折经验分布的单调映射，范围0-6"),
        ("Q1_相对评分", "相对强度等级_整数0至6", "连续相对等级四舍五入并限制在0-6"),
        ("Q1_相对评分", "重复预测标准差", "外层重复折外预测之间的标准差"),
        ("Q2_Top3", "维度内名次", "同一感官维度内的稳健统计排序，1为最高"),
        ("Q2_Top3", "Bootstrap入选频率", "重采样中进入该感官维度Top-3的比例"),
        ("Q2_Top3", "留一风格Top3频率", "逐风格删除后仍进入Top-3的比例"),
        ("Q2_Top3", "化学证据等级", "基于名称映射、方向/OAV及外部核验的A/B/C/D分级"),
        ("Q3_生成画像", "generated_relative_grade_mean", "生成谱经最终感官预测接口得到的平均相对等级"),
        ("Q3_迁移状态", "eligible", "在预先定义统计与直接证据口径下是否可进入正式候选"),
        ("Q3_迁移状态", "physical_safety_certified", "是否已完成物理安全认证；当前均不能据此宣称实际可酿造"),
        ("Q3_生成化学谱", "style/sample_id", "目标风格与该风格内生成样本编号；其余列为231个化学/理化字段"),
    ]
    add_dataframe_sheet(wb, "字段说明", pd.DataFrame(field_rows, columns=["工作表", "字段", "定义或边界"]))

    sources = [q1_path, q1_summary_path, q1_target_path, q2_path, q3_profile_path, q3_migration_path, q3_spectra_path, target_map_path]
    source_rows = []
    for path in sources:
        rows = cols = None
        if path.name.endswith((".csv", ".csv.gz")):
            frame = pd.read_csv(path)
            rows, cols = frame.shape
        source_rows.append({
            "来源文件": str(path.relative_to(ROOT)),
            "行数": rows,
            "列数": cols,
            "SHA256": sha256(path),
        })
    add_dataframe_sheet(wb, "来源校验", pd.DataFrame(source_rows))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    check = load_workbook(out_path, read_only=True, data_only=False)
    expected_sheets = [
        "说明", "Q1_相对评分", "Q1_总体指标", "Q1_逐目标指标", "Q2_Top3",
        "Q3_生成画像", "Q3_迁移状态", "Q3_生成化学谱", "字段说明", "来源校验",
    ]
    if check.sheetnames != expected_sheets:
        raise AssertionError(check.sheetnames)
    row_expectations = {
        "Q1_相对评分": 12501,
        "Q2_Top3": 151,
        "Q3_生成画像": 151,
        "Q3_迁移状态": 34,
        "Q3_生成化学谱": 901,
    }
    for sheet, expected_rows in row_expectations.items():
        if check[sheet].max_row != expected_rows:
            raise AssertionError((sheet, check[sheet].max_row, expected_rows))
    check.close()
    return {
        "path": f"06_交付/02_结果文件/{RESULT_BOOK_NAME}",
        "sha256": sha256(out_path),
        "size_bytes": out_path.stat().st_size,
        "sheet_count": len(expected_sheets),
        "sheets": expected_sheets,
    }


def neutral_copy(src: Path, dst: Path) -> None:
    copy_file(src, dst, sanitize=True)


def copy_support_programs(support: Path) -> None:
    program_root = support / "01_源程序"
    sources = {
        "Q1_感官评分": ROOT / "02_代码/Q1_感官评分",
        "Q2_关键化合物": ROOT / ".worktrees/q2-evidence-opt/02_代码/Q2_关键化合物",
        "Q3_风格生成与迁移": ROOT / ".worktrees/q3-foundation/02_代码/Q3_风格生成",
    }
    for name, source in sources.items():
        copy_tree(source, program_root / name, sanitize=True)

    common = program_root / "公共工具"
    for rel in [
        "scripts/plot_style.R",
        "scripts/session_info.R",
        "scripts/check_result.py",
        "scripts/prepare_submission.py",
    ]:
        copy_file(ROOT / rel, common / Path(rel).name, sanitize=True)

    extra_plots = {
        "Q1_感官评分/补充检验绘图/plot_q1_validation.R": ROOT / "scripts/plot_q1_validation.R",
        "Q1_感官评分/补充检验绘图/plot_q1_validation_s2.R": ROOT / "scripts/plot_q1_validation_s2.R",
        "Q2_关键化合物/论文绘图/plot_q2_paper_v4.R": ROOT / ".worktrees/q2-evidence-opt/04_图表/Q2_关键化合物/v4/plot_q2_paper_v4.R",
        "Q2_关键化合物/补充检验绘图/plot_q2_validation_v2.R": ROOT / "scripts/plot_q2_validation_v2.R",
        "Q3_风格生成与迁移/论文绘图/plot_q3_figures.R": ROOT / ".worktrees/q3-foundation/04_图表/Q3_风格生成/draft_v2/plot_q3_figures.R",
        "Q3_风格生成与迁移/补充检验绘图/plot_q3_validation.R": ROOT / "04_图表/Q3_风格生成/待审_20261027/plot_q3_validation.R",
        "公共工具/plot_total_data_overviews.R": ROOT / "scripts/plot_total_data_overviews.R",
        "公共工具/plot_total_technical_route.R": ROOT / "scripts/plot_total_technical_route.R",
    }
    for rel, src in extra_plots.items():
        copy_file(src, program_root / rel, sanitize=True)

    software_doc = ROOT / "05_论文/交付给论文手/三问附录核心代码_v2/软件与命令说明.md"
    copy_file(software_doc, support / "05_运行环境/软件与命令说明.md", sanitize=True)


def copy_support_data(support: Path) -> None:
    data_root = support / "02_数据_不含赛题原始数据"
    mapping = {
        "01_公共映射/compound_name_map.csv": ROOT / "01_数据/compound_name_map.csv",
        "01_公共映射/感官目标中文与模板列对照.csv": ROOT / "05_论文/总论文/数据与结果_v2/感官目标中文与模板列对照.csv",
        "02_Q1_感官评分/数据契约.json": ROOT / "03_结果/Q1_感官评分/q1_data_contract.json",
        "02_Q1_感官评分/外层交叉验证划分.csv": ROOT / "03_结果/Q1_感官评分/q1_cv_splits.csv",
        "02_Q1_感官评分/外层交叉验证清单.json": ROOT / "03_结果/Q1_感官评分/q1_cv_manifest.json",
        "02_Q1_感官评分/预处理步骤.csv": ROOT / "03_结果/Q1_感官评分/q1_preprocessing_steps.csv",
        "02_Q1_感官评分/特征尺度审计.csv": ROOT / "03_结果/Q1_感官评分/q1_feature_scale_audit.csv",
        "02_Q1_感官评分/缺失结构.csv": ROOT / "03_结果/Q1_感官评分/q1_missingness_by_feature.csv",
        "02_Q1_感官评分/感官目标字典.csv": ROOT / "05_论文/Q1_感官评分/数据与结果/target_dictionary.csv",
        "03_Q2_关键化合物/数据契约.json": ROOT / ".worktrees/q2-evidence-opt/03_结果/Q2_关键化合物/q2_data_contract.json",
        "03_Q2_关键化合物/严格单体候选字典.csv": ROOT / ".worktrees/q2-evidence-opt/03_结果/Q2_关键化合物/q2_candidate_dictionary.csv",
        "03_Q2_关键化合物/感官目标字典.csv": ROOT / ".worktrees/q2-evidence-opt/03_结果/Q2_关键化合物/q2_target_dictionary.csv",
        "04_Q3_风格生成与迁移/字段契约.csv": ROOT / ".worktrees/q3-foundation/03_结果/Q3_风格生成/q3_field_contract_E200.csv",
        "04_Q3_风格生成与迁移/测量状态.csv": ROOT / ".worktrees/q3-foundation/03_结果/Q3_风格生成/q3_measurement_states_E200.csv",
        "04_Q3_风格生成与迁移/预处理步骤.csv": ROOT / ".worktrees/q3-foundation/03_结果/Q3_风格生成/q3_preprocessing_steps_E200.csv",
        "04_Q3_风格生成与迁移/风格样本数.csv": ROOT / ".worktrees/q3-foundation/03_结果/Q3_风格生成/q3_style_counts_E200.csv",
        "04_Q3_风格生成与迁移/缺陷字段登记.csv": ROOT / ".worktrees/q3-foundation/03_结果/Q3_风格生成/q3_defect_registry_E200.csv",
        "04_Q3_风格生成与迁移/源样本缺陷状态.csv": ROOT / ".worktrees/q3-foundation/03_结果/Q3_风格生成/q3_source_defects_E200.csv",
    }
    for rel, src in mapping.items():
        copy_file(src, data_root / rel, sanitize=True)

    readme = """# 派生数据说明

本目录只包含由赛题附件核验或计算得到的字段映射、数据契约、冻结划分、预处理记录和字典。
不包含附件1-8、题面 PDF、承诺书或任何其他官方原始文件。复算前请按根目录运行说明自行放置官方附件。
"""
    (data_root / "README_不含原始数据.md").write_text(readme, encoding="utf-8")


def copy_support_results(support: Path) -> None:
    results = support / "03_结果"

    copy_tree(
        ROOT / "03_结果/Q1_感官评分/publication_final",
        results / "Q1_感官评分/01_最终结果",
        sanitize=True,
    )
    copy_tree(
        ROOT / "05_论文/Q1_感官评分/数据与结果/validation_v1",
        results / "Q1_感官评分/02_补充检验",
        sanitize=True,
    )
    copy_tree(
        ROOT / "05_论文/Q1_感官评分/数据与结果/validation_ablation_v1",
        results / "Q1_感官评分/03_消融与留一风格",
        sanitize=True,
    )

    q2_main = ROOT / "05_论文/Q2_关键化合物/数据与结果"
    for name in [
        "关键化合物Top3_当前版.csv",
        "全部候选稳健排名_当前版.csv",
        "化学证据核验_当前版.csv",
        "模型性能对比_当前版.csv",
    ]:
        copy_file(q2_main / name, results / "Q2_关键化合物/01_最终结果" / name, sanitize=True)
    q2_internal = ROOT / ".worktrees/q2-evidence-opt/03_结果/Q2_关键化合物"
    q2_attr = {
        "最终模型_条件重要性.csv.gz": q2_internal / "q2_final_nnls_conditional_importance_E115.csv.gz",
        "最终模型_模块重要性.csv.gz": q2_internal / "q2_final_nnls_module_importance_E115.csv.gz",
        "最终模型_删除敏感性.csv": q2_internal / "q2_final_nnls_sufficiency_E115.csv",
        "最终模型_归因审计.json": q2_internal / "q2_audit_E115_final_nnls.json",
        "最终模型_归因运行清单.json": q2_internal / "q2_run_manifest_E115_final_nnls.json",
    }
    for name, src in q2_attr.items():
        copy_file(src, results / "Q2_关键化合物/02_独立模型归因" / name, sanitize=True)
    copy_tree(
        ROOT / "05_论文/Q2_关键化合物/数据与结果/validation_v1",
        results / "Q2_关键化合物/03_补充检验",
        sanitize=True,
    )

    total = ROOT / "05_论文/总论文/数据与结果_v2"
    q3_outputs = {
        "三风格_900条生成化学谱.csv.gz": total / "问题三_900条生成化学谱.csv.gz",
        "三风格_完整感官画像.csv": total / "问题三_完整风格画像.csv",
        "跨风格迁移_33组状态.csv": total / "迁移状态全集.csv",
    }
    for name, src in q3_outputs.items():
        copy_file(src, results / "Q3_风格生成与迁移/01_最终结果" / name, sanitize=True)

    q3_internal = ROOT / ".worktrees/q3-foundation/03_结果/Q3_风格生成"
    q3_selected = {
        "主模型_稳健性汇总.csv": q3_internal / "q3_e202_robust_summary.csv",
        "主模型_逐重复指标.csv": q3_internal / "q3_e202_robust_by_repeat.csv",
        "主模型_逐折分布指标.csv": q3_internal / "q3_distribution_metrics_E202.csv",
        "三风格生成画像_逐样本.csv.gz": q3_internal / "q3_generated_profiles_E205.csv.gz",
        "跨风格迁移_候选明细.csv": q3_internal / "q3_migration_summary_E204.csv",
        "跨风格迁移_汇总.csv": q3_internal / "q3_migration_aggregate_E204.csv",
        "跨风格迁移_化学谱.csv.gz": q3_internal / "q3_migration_spectra_E204.csv.gz",
        "跨风格迁移_感官画像.csv.gz": q3_internal / "q3_migration_sensory_profiles_E205.csv.gz",
        "固定主规则候选.csv": q3_internal / "q3_primary_candidates_E207.csv",
        "固定候选裕量敏感性.csv": q3_internal / "q3_frozen_candidate_sensitivity_E207.csv",
    }
    for name, src in q3_selected.items():
        copy_file(src, results / "Q3_风格生成与迁移/02_模型输出与迁移审计" / name, sanitize=True)
    copy_tree(
        ROOT / "05_论文/Q3_风格生成/待审_20261027/数据与结果/validation_v1",
        results / "Q3_风格生成与迁移/03_补充检验",
        sanitize=True,
    )


def safe_filename(text: str) -> str:
    text = text.replace(" ", "_")
    text = re.sub(r"[\\/:*?\"<>|]", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip("_.")


def extract_paper_figures(support: Path) -> list[dict[str, str]]:
    docx = require(ROOT / "05_论文/总论文/2026_MathorCup_大数据竞赛_总论文_v8.docx")
    figure_list = require(ROOT / "05_论文/总论文/总论文_v7_图清单.csv")
    out_dir = support / "04_论文图表/01_论文实际用图"
    out_dir.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, str]] = []
    with figure_list.open("r", encoding="utf-8-sig", newline="") as fh, zipfile.ZipFile(docx) as zf:
        for row in csv.DictReader(fh):
            internal = row["包内图片"]
            data = zf.read(internal)
            actual = hashlib.sha256(data).hexdigest()
            if actual != row["SHA256"]:
                raise AssertionError((row["图号"], actual, row["SHA256"]))
            title = row["图题"]
            clean_title = re.sub(r"^(图\d+|附图S\d+)\s*", "", title).strip()
            if row["图号"].startswith("图"):
                number = int(row["图号"].replace("图", ""))
                filename = f"图{number:02d}_{safe_filename(clean_title)}.png"
            else:
                filename = f"{row['图号']}_{safe_filename(clean_title)}.png"
            dst = out_dir / filename
            dst.write_bytes(data)
            records.append({
                "图号": row["图号"],
                "图题": title,
                "文件": str(dst.relative_to(support)),
                "SHA256": actual,
                "来源": "总论文v8内嵌原图",
            })
    return records


def copy_figure_sources(support: Path) -> None:
    fig_root = support / "04_论文图表/02_可编辑图源"
    selections = {
        "Q1_感官评分/主图": ROOT / "04_图表/Q1_感官评分/final_v4",
        "Q1_感官评分/补充检验": ROOT / "04_图表/Q1_感官评分/validation_v1",
        "Q1_感官评分/绘图数据": ROOT / "03_结果/Q1_感官评分/figure_data",
        "Q2_关键化合物/主图": ROOT / ".worktrees/q2-evidence-opt/04_图表/Q2_关键化合物/v4",
        "Q2_关键化合物/补充检验": ROOT / "04_图表/Q2_关键化合物/待审_20261027",
        "Q3_风格生成与迁移/主图": ROOT / ".worktrees/q3-foundation/04_图表/Q3_风格生成/draft_v2",
        "Q3_风格生成与迁移/补充检验": ROOT / "04_图表/Q3_风格生成/待审_20261027",
        "总论文/技术路线": ROOT / "05_论文/总论文/图稿_待审计/技术路线",
        "总论文/数据概览": ROOT / "05_论文/总论文/图稿_待审计/数据概览",
    }
    for name, src in selections.items():
        copy_tree(src, fig_root / name, sanitize=True)


def classify_program(path: Path, support: Path) -> tuple[str, str, str]:
    rel = path.relative_to(support).as_posix()
    low = rel.lower()
    name_low = path.name.lower()
    if path.suffix.lower() == ".md" and ("readme" in name_low or "说明" in path.name):
        category = "说明"
    elif "/tests/" in low or Path(rel).name.lower().startswith("test_"):
        category = "测试"
    elif "/configs/" in low or path.suffix.lower() in {".yaml", ".yml"}:
        category = "配置"
    elif any(k in low for k in ["plot", "figure", "绘图"]):
        category = "绘图与图源准备"
    elif any(k in low for k in ["contract", "preprocess", "split", "dictionary", "prepare_abv"]):
        category = "数据契约与预处理"
    elif any(k in low for k in ["audit", "verify", "metric", "evaluation", "diagnostic", "compare", "summarize"]):
        category = "评价与审计"
    elif any(k in low for k in ["docx", "handoff", "omml"]):
        category = "论文生成辅助"
    else:
        category = "模型与求解"

    core_tokens = {
        "Q1": ["data_contract.py", "make_cv_splits.py", "baseline_models.py", "ensemble_models.py", "nnls_stacking.py", "q1_metrics.py", "run_e009_nnls_stacking.py", "plot_q1_final.r"],
        "Q2": ["q2_contract.py", "q2_pipeline.py", "rank_consensus.py", "robust_consensus.py", "bootstrap_stability.py", "chemical_evidence.py", "q1_final_adapter.py", "hierarchical_importance.py", "run_e111_robust_consensus.py", "run_e112_chemical_evidence.py", "run_e115_final_nnls.py", "plot_q2_paper_v4.r"],
        "Q3": ["q3_contract.py", "q3_constraints.py", "q3_evaluation.py", "gaussian_baselines.py", "constrained_migration.py", "sensory_audit.py", "q1_interface.py", "run_e202_gaussian_baselines.py", "run_e204_constrained_migration.py", "run_e205_sensory_audit.py", "plot_q3_figures.r"],
    }
    question = "公共"
    if "q1_感官评分" in low:
        question = "Q1"
    elif "q2_关键化合物" in low:
        question = "Q2"
    elif "q3_风格生成" in low:
        question = "Q3"
    is_core = "是" if question in core_tokens and any(token in low for token in core_tokens[question]) else "否"
    return question, category, is_core


def compound_suffix(path: Path) -> tuple[str, str]:
    """Return a filename stem and a compound suffix such as .csv.gz."""
    suffix = "".join(path.suffixes)
    if not suffix:
        return path.name, ""
    return path.name[: -len(suffix)], suffix


def short_component(text: str, *, limit: int = 92) -> str:
    text = safe_filename(text)
    if len(text) <= limit:
        return text
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()[:8]
    return f"{text[: limit - 9]}_{digest}"


def section_metadata(section: str, rel: Path, path: Path, support: Path) -> tuple[str, str, str, str]:
    """Return question, sortable category code, category label, and core flag."""
    rel_text = rel.as_posix()
    rel_low = rel_text.lower()
    question = "公共"
    prefix = "COM"
    if "q1_感官评分" in rel_low:
        question, prefix = "Q1", "Q1"
    elif "q2_关键化合物" in rel_low:
        question, prefix = "Q2", "Q2"
    elif "q3_风格生成" in rel_low:
        question, prefix = "Q3", "Q3"

    if section == "01_源程序":
        question, label, core = classify_program(path, support)
        prefix = {"Q1": "Q1", "Q2": "Q2", "Q3": "Q3", "公共": "COM"}[question]
        order = {
            "说明": "00",
            "配置": "10",
            "数据契约与预处理": "20",
            "模型与求解": "30",
            "评价与审计": "40",
            "绘图与图源准备": "50",
            "测试": "60",
            "论文生成辅助": "70",
        }
        return question, f"{order[label]}_{label.replace('与图源准备', '')}", label, core

    if section == "02_数据_不含赛题原始数据":
        if rel.parts and rel.parts[0].startswith("01_"):
            question, prefix = "公共", "COM"
        name = path.name.lower()
        if path.suffix.lower() == ".md":
            code, label = "00", "说明"
        elif any(k in name for k in ["map", "对照", "映射"]):
            code, label = "10", "字段映射"
        elif any(k in name for k in ["contract", "契约"]):
            code, label = "20", "数据契约"
        elif any(k in name for k in ["split", "划分"]):
            code, label = "30", "冻结划分"
        elif any(k in name for k in ["preprocess", "预处理", "缺失", "尺度"]):
            code, label = "40", "预处理审计"
        elif any(k in name for k in ["dictionary", "字典", "登记"]):
            code, label = "50", "字段字典"
        else:
            code, label = "60", "派生数据"
        return question, f"{code}_{label}", label, "否"

    if section == "03_结果":
        folder = rel.parts[1] if len(rel.parts) > 1 else ""
        if folder.startswith("01_"):
            code, label = "10", "正式结果"
        elif folder.startswith("02_"):
            code, label = "20", "模型解释与审计"
        elif folder.startswith("03_"):
            code, label = "30", "补充检验"
        else:
            code, label = "40", "其他结果"
        return question, f"{code}_{label}", label, "否"

    if section == "04_论文图表":
        if rel.parts and rel.parts[0] == "01_论文实际用图":
            return "总论文", "10_正式用图", "正式用图", "否"
        if any(part.lower() == "qa" for part in rel.parts) or any(
            token in rel_low for token in ["gray", "灰度", "70pct", "contact", "审图合页", "彩色灰度检查"]
        ):
            label_code = "90_质量检查"
        elif "绘图数据" in rel_text or "source_data" in rel_low:
            label_code = "40_绘图数据"
        elif "补充检验" in rel_text:
            label_code = "30_补充图源"
        elif "主图" in rel_text:
            label_code = "20_主图图源"
        elif "技术路线" in rel_text:
            label_code = "50_技术路线"
        elif "数据概览" in rel_text:
            label_code = "60_数据概览"
        else:
            label_code = "70_其他图源"
        label = label_code.split("_", 1)[1]
        if "总论文" in rel_text:
            question, prefix = "总论文", "DOC"
        return question, label_code, label, "否"

    raise ValueError(section)


def flatten_support_sections(support: Path) -> list[dict[str, object]]:
    """Flatten the four user-facing sections and return a reversible mapping."""
    sections = ["01_源程序", "02_数据_不含赛题原始数据", "03_结果", "04_论文图表"]
    records: list[dict[str, object]] = []
    for section in sections:
        source_root = require(support / section)
        files = [p for p in sorted(source_root.rglob("*")) if p.is_file()]
        temp_root = support / f".__flat_{section}"
        temp_root.mkdir(parents=True, exist_ok=False)
        counters: dict[tuple[str, str], int] = {}
        used_names: set[str] = set()
        for path in files:
            rel = path.relative_to(source_root)
            question, category_code, category_label, is_core = section_metadata(section, rel, path, support)
            prefix = {"Q1": "Q1", "Q2": "Q2", "Q3": "Q3", "公共": "COM", "总论文": "DOC"}[question]
            key = (prefix, category_code)
            counters[key] = counters.get(key, 0) + 1
            sequence = counters[key]
            stem, suffix = compound_suffix(path)
            stem = short_component(stem)
            new_name = f"{prefix}_{category_code}_{sequence:03d}_{stem}{suffix}"
            if new_name in used_names:
                digest = hashlib.sha256(rel.as_posix().encode("utf-8")).hexdigest()[:8]
                new_name = f"{prefix}_{category_code}_{sequence:03d}_{stem}_{digest}{suffix}"
            if new_name in used_names:
                raise AssertionError(f"Flatten collision: {new_name}")
            used_names.add(new_name)
            dst = temp_root / new_name
            shutil.copy2(path, dst)
            records.append({
                "区段": section,
                "问题": question,
                "类别": category_label,
                "是否核心": is_core,
                "排序前缀": f"{prefix}_{category_code}_{sequence:03d}",
                "原目录路径": rel.as_posix(),
                "平铺文件名": new_name,
                "SHA256": sha256(dst),
            })
        shutil.rmtree(source_root)
        temp_root.rename(source_root)

    manifest_root = support / "06_清单与校验"
    write_csv(
        manifest_root / "平铺文件映射.csv",
        ["区段", "问题", "类别", "是否核心", "排序前缀", "原目录路径", "平铺文件名", "SHA256"],
        records,
    )
    return records


def remap_figure_records(
    figure_records: list[dict[str, str]], flat_records: list[dict[str, object]]
) -> list[dict[str, str]]:
    lookup = {
        f"{row['区段']}/{row['原目录路径']}": f"{row['区段']}/{row['平铺文件名']}"
        for row in flat_records
    }
    remapped = []
    for row in figure_records:
        old = row["文件"]
        if old not in lookup:
            raise AssertionError(f"Figure missing from flatten map: {old}")
        updated = dict(row)
        updated["文件"] = lookup[old]
        remapped.append(updated)
    return remapped


def write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def build_manifests(
    support: Path,
    figure_records: list[dict[str, str]],
    flat_records: list[dict[str, object]],
) -> dict[str, object]:
    manifest_root = support / "06_清单与校验"
    manifest_root.mkdir(parents=True, exist_ok=True)

    program_rows = []
    for record in [r for r in flat_records if r["区段"] == "01_源程序"]:
        path = support / "01_源程序" / str(record["平铺文件名"])
        program_rows.append({
            "问题": record["问题"],
            "类别": record["类别"],
            "是否核心": record["是否核心"],
            "原目录路径": record["原目录路径"],
            "文件": str(path.relative_to(support)),
            "大小_字节": path.stat().st_size,
            "SHA256": sha256(path),
        })
    write_csv(
        manifest_root / "程序清单.csv",
        ["问题", "类别", "是否核心", "原目录路径", "文件", "大小_字节", "SHA256"],
        program_rows,
    )

    data_rows = []
    for record in [r for r in flat_records if r["区段"] == "02_数据_不含赛题原始数据"]:
        path = support / "02_数据_不含赛题原始数据" / str(record["平铺文件名"])
        data_rows.append({
            "问题": record["问题"],
            "类别": record["类别"],
            "原目录路径": record["原目录路径"],
            "文件": str(path.relative_to(support)),
            "大小_字节": path.stat().st_size,
            "SHA256": sha256(path),
        })
    write_csv(
        manifest_root / "数据清单.csv",
        ["问题", "类别", "原目录路径", "文件", "大小_字节", "SHA256"],
        data_rows,
    )

    result_rows = []
    for record in [r for r in flat_records if r["区段"] == "03_结果"]:
        path = support / "03_结果" / str(record["平铺文件名"])
        result_rows.append({
            "问题": record["问题"],
            "类别": record["类别"],
            "原目录路径": record["原目录路径"],
            "文件": str(path.relative_to(support)),
            "大小_字节": path.stat().st_size,
            "SHA256": sha256(path),
        })
    write_csv(
        manifest_root / "结果清单.csv",
        ["问题", "类别", "原目录路径", "文件", "大小_字节", "SHA256"],
        result_rows,
    )

    figure_source_rows = []
    for record in [r for r in flat_records if r["区段"] == "04_论文图表"]:
        path = support / "04_论文图表" / str(record["平铺文件名"])
        figure_source_rows.append({
            "问题": record["问题"],
            "类别": record["类别"],
            "原目录路径": record["原目录路径"],
            "文件": str(path.relative_to(support)),
            "大小_字节": path.stat().st_size,
            "SHA256": sha256(path),
        })
    write_csv(
        manifest_root / "图源文件清单.csv",
        ["问题", "类别", "原目录路径", "文件", "大小_字节", "SHA256"],
        figure_source_rows,
    )
    write_csv(manifest_root / "论文用图清单.csv", ["图号", "图题", "文件", "SHA256", "来源"], figure_records)

    exclusions = """# 明确排除项

- 未打包赛题附件1-8、题面 PDF、论文模板、承诺书及其他官方原始文件；
- 未打包 `__pycache__`、`.pyc`、pytest 缓存、Git 元数据、日志、临时文件；
- 未打包模型 checkpoint、joblib/pickle 缓存、神经网络权重或大体积搜索缓存；
- 未把未采用方案改写为正式结论；问题二正式排序和问题三主模型口径保持冻结；
- 论文实际用图按总论文 v8 内嵌字节提取，可编辑图源另列。
"""
    (manifest_root / "排除项说明.md").write_text(exclusions, encoding="utf-8")

    all_files = [p for p in sorted(support.rglob("*")) if p.is_file() and p.name != "SHA256SUMS.txt"]
    checksum_path = manifest_root / "SHA256SUMS.txt"
    with checksum_path.open("w", encoding="utf-8") as fh:
        for path in all_files:
            fh.write(f"{sha256(path)}  {path.relative_to(support).as_posix()}\n")

    counts = {
        "program_files": len(program_rows),
        "data_files": len(data_rows),
        "result_files": len(result_rows),
        "figure_source_files": len(figure_source_rows),
        "paper_figures": len(figure_records),
        "initial_checksum_entries": len(all_files),
    }
    (manifest_root / "交付统计.json").write_text(
        json.dumps(counts, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return counts


def write_support_docs(support: Path) -> None:
    readme = """# A赛道支撑材料说明

本目录按“源程序-派生数据-结果-论文图表-运行环境-校验清单”分层整理。按提交要求，前四个目录内部均为单层平铺；文件名以前缀 `Q1/Q2/Q3/COM/DOC + 类别序号 + 文件序号` 排序。

## 1. 正式口径

- Q1：感官组收缩非负最小二乘堆叠，10×3 外层重复折外评价，输出50维标准化预测与0-6相对强度等级；
- Q2：稳健统计 Top-3 为正式答案，最终 Q1 模型归因仅作独立解释证据，不改变排序；
- Q3：经验边缘与风格收缩 Gaussian Copula 为生成主模型，迁移采用显式缺陷约束与证据状态审计。

## 2. 目录

- `01_源程序/`：三问完整计算工程平铺文件；名称中的 `00/10/20/30/40/50/60/70` 分别表示说明、配置、数据处理、模型求解、评价审计、绘图、测试和论文辅助；
- `02_数据_不含赛题原始数据/`：平铺的字段映射、契约、冻结划分与派生字典；
- `03_结果/`：平铺的正式结果、解释审计与补充检验；
- `04_论文图表/`：平铺的总论文 v8 实际使用21张图，以及可编辑图源、绘图数据和质量检查文件；
- `05_运行环境/`：软件版本、依赖与命令说明；
- `06_清单与校验/`：程序、数据、结果、图表清单、旧路径到平铺文件名映射与 SHA256。

## 3. 复算准备

本目录不含赛题原始附件。平铺文件用于提交与审阅；如需按原工程路径复算，先运行 `python 05_运行环境/还原平铺文件.py --output 复算工程`，该命令依据映射表复制还原四个目录，不改变平铺交付文件。随后补放官方附件并按各问运行说明执行。R 图统一调用 `theme_mcup()`。

## 4. 使用边界

工作簿中的 Q1 0-6 数值是相对强度映射；Q2 为稳定关键预测关联而非因果效应；Q3 为统计生成与迁移设计，不能直接等同于实际可酿造或物理安全认证。
"""
    (support / "README_支撑材料说明.md").write_text(readme, encoding="utf-8")

    per_question = {
        "Q1_感官评分": """# Q1 运行说明

本说明文件在平铺后带 `Q1_00_说明` 前缀。复算前先按根说明还原目录；核心入口为数据契约、冻结划分、非负堆叠求解和最终绘图程序。

```bash
conda activate mathorcup
cd 复算工程
python 01_源程序/Q1_感官评分/src/data_contract.py
python 01_源程序/Q1_感官评分/src/make_cv_splits.py
python 01_源程序/Q1_感官评分/src/run_e009_nnls_stacking.py --stage full
Rscript 01_源程序/Q1_感官评分/src/plot_q1_final.R
```
""",
        "Q2_关键化合物": """# Q2 运行说明

本说明文件在平铺后带 `Q2_00_说明` 前缀。复算前先按根说明还原目录；核心链为数据契约、四路排序与稳健融合、Bootstrap/留一风格稳定性、化学证据分级，最终模型归因只作独立解释。

```bash
conda activate mathorcup
cd 复算工程
python 01_源程序/Q2_关键化合物/src/run_e100_contract.py
python 01_源程序/Q2_关键化合物/src/run_e101_e104.py --repeats 10 --run-label ''
python 01_源程序/Q2_关键化合物/src/run_e107_bootstrap.py --n-bootstraps 500 --run-label ''
python 01_源程序/Q2_关键化合物/src/run_e111_robust_consensus.py
python 01_源程序/Q2_关键化合物/src/run_e112_chemical_evidence.py
python 01_源程序/Q2_关键化合物/src/run_e115_final_nnls.py
Rscript 01_源程序/Q2_关键化合物/论文绘图/plot_q2_paper_v4.R
```
""",
        "Q3_风格生成与迁移": """# Q3 运行说明

本说明文件在平铺后带 `Q3_00_说明` 前缀。复算前先按根说明还原目录；核心入口为数据契约、风格收缩 Copula、受约束迁移及最终感官接口，其他生成结构作为消融保留。

```bash
conda activate mathorcup
cd 复算工程
python 01_源程序/Q3_风格生成与迁移/src/run_e200_contract.py --write
python 01_源程序/Q3_风格生成与迁移/src/run_e202_gaussian_baselines.py
python 01_源程序/Q3_风格生成与迁移/src/run_e204_constrained_migration.py
python 01_源程序/Q3_风格生成与迁移/src/run_e205_sensory_audit.py
Rscript 01_源程序/Q3_风格生成与迁移/论文绘图/plot_q3_figures.R
```
""",
    }
    for name, text in per_question.items():
        (support / "01_源程序" / name / "交付运行说明.md").write_text(text, encoding="utf-8")


def copy_environment(support: Path) -> None:
    env_dir = support / "05_运行环境"
    for rel in ["requirements.txt", "docs/env_py.txt", "docs/env_r.txt"]:
        copy_file(ROOT / rel, env_dir / Path(rel).name, sanitize=True)
    versions = []
    commands = [
        [sys.executable, "--version"],
        [shutil.which("Rscript") or str(Path("/home/mary/miniconda3/envs/r/bin/Rscript")), "--version"],
    ]
    for cmd in commands:
        try:
            proc = subprocess.run(cmd, cwd=ROOT, text=True, capture_output=True, timeout=30, check=False)
            versions.append(
                sanitize_text(f"$ {' '.join(cmd)}\n{(proc.stdout + proc.stderr).strip()}\n")
            )
        except Exception as exc:  # pragma: no cover - environment-dependent note
            versions.append(f"$ {' '.join(cmd)}\nERROR: {exc}\n")
    (env_dir / "版本核验.txt").write_text("\n".join(versions), encoding="utf-8")

    restore_script = '''#!/usr/bin/env python3
"""Restore the four flattened support sections into their original layout."""
from __future__ import annotations
import argparse
import csv
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MAPPING = ROOT / "06_清单与校验/平铺文件映射.csv"

def main() -> None:
    parser = argparse.ArgumentParser(description="还原平铺的源程序、派生数据、结果与图源目录")
    parser.add_argument("--output", required=True, help="新建的还原目录；必须不存在或为空")
    args = parser.parse_args()
    output = Path(args.output).resolve()
    if output.exists() and any(output.iterdir()):
        raise SystemExit(f"输出目录非空，停止：{output}")
    output.mkdir(parents=True, exist_ok=True)
    with MAPPING.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    for row in rows:
        src = ROOT / row["区段"] / row["平铺文件名"]
        dst = output / row["区段"] / row["原目录路径"]
        if not src.is_file():
            raise FileNotFoundError(src)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
    print(f"已还原 {len(rows)} 个文件到 {output}")

if __name__ == "__main__":
    main()
'''
    (env_dir / "还原平铺文件.py").write_text(restore_script, encoding="utf-8")

    flat_note = """# 平铺目录使用说明

`01_源程序`、`02_数据_不含赛题原始数据`、`03_结果`、`04_论文图表` 均为单层目录。
文件名前缀依次表示问题、类别和顺序；完整旧路径见 `06_清单与校验/平铺文件映射.csv`。

如需恢复原来的可运行目录结构，在支撑材料根目录执行：

```bash
python 05_运行环境/还原平铺文件.py --output 复算工程
```

还原程序只复制文件，不修改或删除平铺交付文件。官方原始附件未打包，完整复算前仍需自行放回约定数据目录。
"""
    (env_dir / "平铺目录使用说明.md").write_text(flat_note, encoding="utf-8")


def audit_delivery(result_dir: Path, support: Path) -> dict[str, object]:
    result_files = [p for p in result_dir.iterdir() if p.is_file()]
    if [p.name for p in result_files] != [RESULT_BOOK_NAME]:
        raise AssertionError(f"Result directory must contain one workbook only: {result_files}")
    if any(p.suffix.lower() != ".xlsx" for p in result_files):
        raise AssertionError(result_files)

    forbidden_parts = {"data", "2026初赛", "附件", "__pycache__", ".git", ".pytest_cache"}
    violations = []
    absolute_hits = []
    flat_section_violations = []
    for section in ["01_源程序", "02_数据_不含赛题原始数据", "03_结果", "04_论文图表"]:
        section_root = support / section
        nested = [p for p in section_root.rglob("*") if p.is_dir()]
        flat_section_violations.extend(str(p.relative_to(support)) for p in nested)
    if flat_section_violations:
        raise AssertionError(f"Sections are not flat: {flat_section_violations[:20]}")
    for path in support.rglob("*"):
        if not path.is_file():
            continue
        if any(part in forbidden_parts for part in path.relative_to(support).parts):
            violations.append(str(path.relative_to(support)))
        if path.suffix in TEXT_EXTENSIONS:
            try:
                text = path.read_text(encoding="utf-8-sig")
            except UnicodeDecodeError:
                continue
            if any(token in text for token in ["/home/mary", "/mnt/d/上海大学", r"D:\上海大学"]):
                absolute_hits.append(str(path.relative_to(support)))
    if violations:
        raise AssertionError(f"Forbidden paths in support: {violations[:20]}")
    if absolute_hits:
        raise AssertionError(f"Unsanitized local paths: {absolute_hits[:20]}")
    if any(p.name.endswith((".pyc", ".pyo", ".joblib", ".pkl", ".pickle")) for p in support.rglob("*")):
        raise AssertionError("Cache or bytecode leaked into support materials")

    check = load_workbook(result_files[0], read_only=True, data_only=False)
    workbook_rows = {name: check[name].max_row for name in check.sheetnames}
    check.close()
    return {
        "result_file_count": 1,
        "result_workbook_sha256": sha256(result_files[0]),
        "result_workbook_size_bytes": result_files[0].stat().st_size,
        "workbook_rows": workbook_rows,
        "support_file_count": sum(1 for p in support.rglob("*") if p.is_file()),
        "support_size_bytes": sum(p.stat().st_size for p in support.rglob("*") if p.is_file()),
        "forbidden_path_hits": len(violations),
        "absolute_local_path_hits": len(absolute_hits),
        "raw_attachment_files": 0,
        "flat_section_nested_directories": len(flat_section_violations),
    }


def main() -> None:
    DELIVERY_ROOT.mkdir(parents=True, exist_ok=True)
    staging_base = Path(tempfile.mkdtemp(prefix="delivery_build_", dir=DELIVERY_ROOT))
    result_staging = staging_base / "02_结果文件"
    support_staging = staging_base / "04_支撑材料"
    result_staging.mkdir(parents=True)
    support_staging.mkdir(parents=True)

    workbook_info = build_result_workbook(result_staging / RESULT_BOOK_NAME)
    copy_support_programs(support_staging)
    copy_support_data(support_staging)
    copy_support_results(support_staging)
    figure_records = extract_paper_figures(support_staging)
    copy_figure_sources(support_staging)
    copy_environment(support_staging)
    write_support_docs(support_staging)
    flat_records = flatten_support_sections(support_staging)
    figure_records = remap_figure_records(figure_records, flat_records)
    support_counts = build_manifests(support_staging, figure_records, flat_records)
    preinstall_audit = audit_delivery(result_staging, support_staging)

    report = {
        "built_at": datetime.now().isoformat(timespec="seconds"),
        "workbook": workbook_info,
        "support": support_counts,
        "audit": preinstall_audit,
    }
    (support_staging / "06_清单与校验/最终验收.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    # Refresh SHA256SUMS after the final audit file is added.
    checksum_path = support_staging / "06_清单与校验/SHA256SUMS.txt"
    with checksum_path.open("w", encoding="utf-8") as fh:
        for path in [p for p in sorted(support_staging.rglob("*")) if p.is_file() and p != checksum_path]:
            fh.write(f"{sha256(path)}  {path.relative_to(support_staging).as_posix()}\n")

    # Re-audit the exact staging trees that will be installed.
    final_audit = audit_delivery(result_staging, support_staging)
    report["audit"] = final_audit
    (support_staging / "06_清单与校验/最终验收.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with checksum_path.open("w", encoding="utf-8") as fh:
        for path in [p for p in sorted(support_staging.rglob("*")) if p.is_file() and p != checksum_path]:
            fh.write(f"{sha256(path)}  {path.relative_to(support_staging).as_posix()}\n")

    # Move the two completed trees out of the common staging parent first.
    result_ready = DELIVERY_ROOT / f".ready_results_{TIMESTAMP}"
    support_ready = DELIVERY_ROOT / f".ready_support_{TIMESTAMP}"
    result_staging.rename(result_ready)
    support_staging.rename(support_ready)
    staging_base.rmdir()

    result_backup = install_directory(result_ready, RESULT_TARGET)
    support_backup = install_directory(support_ready, SUPPORT_TARGET)
    # User-approved delivery format: the final public source folder contains
    # only the curated model/solver/plotting Python and R files.
    curate_script = ROOT / "scripts/curate_core_submission.py"
    if curate_script.is_file():
        subprocess.run(
            [sys.executable, str(curate_script)],
            cwd=ROOT, check=True,
        )
    installed_audit = audit_delivery(RESULT_TARGET, SUPPORT_TARGET)
    print(json.dumps({
        "result": str(RESULT_TARGET / RESULT_BOOK_NAME),
        "support": str(SUPPORT_TARGET),
        "result_backup": str(result_backup) if result_backup else None,
        "support_backup": str(support_backup) if support_backup else None,
        "audit": installed_audit,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
