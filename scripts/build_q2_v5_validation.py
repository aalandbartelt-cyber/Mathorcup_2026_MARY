#!/usr/bin/env python3
"""Build Q2 paper v5: v4 + fixed figure 2 + validation section/figure 5/table 5."""
from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree
from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path.home() / "mathorcup/.worktrees/q2-evidence-opt"
PAPER = ROOT / "05_论文/Q2_关键化合物"
FIG = ROOT / "04_图表/Q2_关键化合物/v4"
VAL = ROOT / "04_图表/Q2_关键化合物/validation_v1"
SOURCE = PAPER / "第二问_论文正文_v4.docx"
TARGET = PAPER / "第二问_论文正文_v5.docx"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def replace_image_media():
    """Swap the four embedded v4 figures, including the fixed figure 2."""
    with ZipFile(SOURCE) as archive:
        blobs = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    root = etree.fromstring(blobs["word/document.xml"])
    rels = etree.fromstring(blobs["word/_rels/document.xml.rels"])
    relmap = {x.get("Id"): x.get("Target") for x in rels}
    drawings = root.xpath(".//w:drawing", namespaces=NS)
    figures = [next(FIG.glob(f"图{i}_*_v4.png")) for i in range(1, 5)]
    assert len(drawings) == 4, len(drawings)
    for drawing, figure in zip(drawings, figures):
        blip = drawing.find(".//a:blip", NS)
        member = "word/" + relmap[blip.get(f"{{{NS['r']}}}embed")]
        blobs[member] = figure.read_bytes()
        width, height = Image.open(figure).size
        extent = drawing.find(".//wp:extent", NS)
        cy = round(int(extent.get("cx")) * height / width)
        extent.set("cy", str(cy))
        for node in drawing.findall(".//a:xfrm/a:ext", NS):
            node.set("cy", str(cy))
    blobs["word/document.xml"] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    with ZipFile(TARGET, "w", ZIP_DEFLATED) as archive:
        for name, blob in blobs.items():
            archive.writestr(name, blob)


def replace_text(doc: Document, old: str, new: str) -> bool:
    for p in doc.paragraphs:
        if old not in p.text:
            continue
        for run in p.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                return True
        # Fallback: text spans runs; preserve paragraph style but not intra-run formatting.
        full = p.text.replace(old, new)
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = full
        else:
            p.add_run(full)
        return True
    return False


def set_cell(cell, value: str, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(value)
    run.bold = bold
    run.font.size = Pt(8.5)
    run.font.name = "SimSun"


def add_validation(doc: Document):
    target = next(p for p in doc.paragraphs if p.text.strip() == "合理性与局限")
    fig4_caption = next(p for p in doc.paragraphs if p.text.strip().startswith("图4 "))
    caption_style = fig4_caption.style

    def add_before(text: str, style):
        return target.insert_paragraph_before(text, style=style)

    add_before("补充检验：负对照、直接锚点与删除敏感性", "Heading 3")
    add_before(
        "为检验排名是否只是候选字段筛选的结果，补充目标置换、影子字段和直接锚点检查。"
        "四个预登记入选阈值下，目标置换均未产生入选关系；影子字段仍有1,524/2,000次局部Top-3入选，"
        "真实字段为4,476/2,000次。因此只能表述为真实字段整体占优，不能写成零误选。",
        "Body Text",
    )
    add_before(
        "直接锚点方面，现行最终Top-3保留4-VG、酒精感、酒体和碳酸感四个既有锚点；"
        "DMS、双乙酰和反-2-壬烯醛分别位于第29、92和61名，只作探索性结论。"
        "删除Top-3后，50个目标的平均ΔMAE为0.00058，22/50个目标出现平均退化；"
        "最大退化为酒精感+0.02244（+6.99%），说明排序整体稳定但不是不可替代。",
        "Body Text",
    )
    add_before(
        "稳定性口径需分开报告：分层自助候选集的入选频率中位数为0.428（43/150不低于0.60），"
        "四路融合最终Top-3为0.395（40/150不低于0.60）；留一风格平均保留率分别为0.641和0.608，"
        "最终Top-3的中位保留率为0.659。正式答案采用四路融合最终Top-3；0.428/43只描述自助候选集的稳定性。",
        "Body Text",
    )

    pic_p = add_before("", "Body Text")
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(VAL / "图5_负对照与直接锚点.png"), width=Cm(15.5))
    cap = add_before("图5 负对照、影子字段与直接锚点恢复", caption_style)
    cap.alignment = fig4_caption.alignment

    tcap = add_before("表5 补充检验的数值结果", caption_style)
    tcap.alignment = fig4_caption.alignment

    rows = [
        ("检验项", "结果", "口径"),
        ("目标置换", "0/0/0/0", "阈值0.60/0.70/0.80/0.90，风格内置换"),
        ("影子字段", "1,524 次；真实字段 4,476 次", "2,000条目标—重复—半样本记录"),
        ("现行直接锚点", "4/4；CO2 bootstrap第4名", "现行最终Top-3；E118候选遗漏CO2"),
        ("删除Top-3", "平均ΔMAE 0.00058；22/50目标变差", "50目标，ΔMAE=masked−baseline"),
        ("稳定性口径", "候选集0.428/43；最终Top-3 0.395/40", "频率中位数与≥0.60条数"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.0), Cm(6.8), Cm(5.6)]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell(table.cell(i, j), value, bold=(i == 0))
            table.cell(i, j).width = widths[j]
    target._p.addprevious(table._tbl)

    add_before(
        "注：影子字段入选不能解释为零误选；删除敏感性为固定最终模型下的Top-3遮蔽，不包含重训不确定性。"
        "DMS、双乙酰和反-2-壬烯醛未进入Top-3，不能写成已恢复的缺陷锚点。",
        "Body Text",
    )


def main():
    replace_image_media()
    doc = Document(TARGET)
    replacements = [
        ("自助入选频率中位数为0.428", "分层自助候选集的入选频率中位数为0.428，四路融合最终Top-3为0.395"),
        ("中位数为0.428；达到0.60", "中位数为0.428；四路融合最终Top-3的中位数为0.395；达到0.60"),
        ("当前150条候选在22次删除中的平均保留率为0.641",
         "分层自助候选集的平均保留率为0.641；四路融合最终Top-3为0.608（中位0.659）"),
        ("图ElasticNet指有符号图正则回归", "图正则EN指有符号图正则回归"),
    ]
    changed = []
    for old, new in replacements:
        if replace_text(doc, old, new):
            changed.append(old)
    add_validation(doc)
    doc.save(TARGET)
    # Keep the companion documents on the same version.
    for stem in ("模型假设_Q2", "问题分析_Q2"):
        src = PAPER / f"{stem}_v4.docx"
        dst = PAPER / f"{stem}_v5.docx"
        if src.exists():
            dst.write_bytes(src.read_bytes())
    audit = {
        "source": str(SOURCE.relative_to(ROOT)),
        "target": str(TARGET.relative_to(ROOT)),
        "changed_text": changed,
        "added": ["补充检验 Heading 3", "图5", "表5"],
        "figures": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in [*[next(FIG.glob(f"图{i}_*_v4.png")) for i in range(1, 5)],
                              VAL / "图5_负对照与直接锚点.png"]},
    }
    (ROOT / "03_结果/Q2_关键化合物/publication_audit/q2_v5_changes.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(TARGET)
    print("changed", changed)


if __name__ == "__main__":
    main()
